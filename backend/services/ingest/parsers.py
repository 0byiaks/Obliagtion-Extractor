# backend/services/ingest/parsers.py
from io import BytesIO
from typing import Tuple
from docx import Document

SUPPORTED_EXTS = {".docx"}

class IngestError(Exception):
    pass

def _clean_lines(lines):
    # Trim and drop empty runs created by formatting artifacts
    out = []
    for ln in lines:
        s = " ".join(ln.split())  # collapse internal whitespace
        if s.strip():
            out.append(s.strip())
    return out

def _docx_to_text(file_bytes: bytes) -> str:
    bio = BytesIO(file_bytes)
    doc = Document(bio)

    lines = []

    # 1) headers (optional; may be empty)
    try:
        for section in doc.sections:
            header = section.header
            if header:
                for p in header.paragraphs:
                    if p.text.strip():
                        lines.append(p.text)
    except Exception:
        # headers not critical — ignore on failure
        pass

    # 2) main body paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)

    # 3) tables (row by row, cell by cell)
    for tbl in doc.tables:
        for row in tbl.rows:
            row_text = []
            for cell in row.cells:
                cell_txt = " ".join([p.text for p in cell.paragraphs if p.text.strip()])
                if cell_txt.strip():
                    row_text.append(cell_txt.strip())
            if row_text:
                lines.append(" | ".join(row_text))

    # 4) footers (optional)
    try:
        for section in doc.sections:
            footer = section.footer
            if footer:
                for p in footer.paragraphs:
                    if p.text.strip():
                        lines.append(p.text)
    except Exception:
        pass

    lines = _clean_lines(lines)
    # Join with newlines to keep light structure; later stages can pre-chunk by lines/blank lines
    return "\n".join(lines)

def to_text(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Returns (text, detected_type). For now, only .docx is supported.
    Raise IngestError on unsupported or parsing failure.
    """
    lname = filename.lower().strip()
    ext = ""
    for e in SUPPORTED_EXTS:
        if lname.endswith(e):
            ext = e
            break

    if ext == ".docx":
        try:
            text = _docx_to_text(file_bytes)
            return text, "docx"
        except Exception as e:
            raise IngestError(f"Failed to parse .docx: {e}")
    else:
        raise IngestError(f"Unsupported file type for demo: {filename}. Supported: .docx")
